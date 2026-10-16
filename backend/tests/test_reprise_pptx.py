"""Un échec en cours de traduction PPTX ne doit pas jeter ce qui est déjà traduit.

LE DÉFAUT QUE CETTE SUITE VERROUILLE
------------------------------------
Les diapositives traduites ne vivaient qu'en mémoire jusqu'à la fin du job :
un échec à la dernière (disque plein, fournisseur coupé) perdait toutes les
autres, déjà payées, et la relance repartait de zéro. Et l'original, écrit
disque plein, restait à 0 octet sans jamais être réécrit : le document
devenait irrécupérable.

On rejoue le VRAI job sur un PPTX synthétique, avec un traducteur factice
(aucun réseau) et la conversion LibreOffice neutralisée (l'aperçu est
best-effort, il n'est pas l'objet de cette suite).

    backend/venv/Scripts/python.exe backend/tests/test_reprise_pptx.py
"""
from __future__ import annotations

import json
import os
import shutil
import sys
import tempfile

import racine  # noqa: F401  -- met backend/ sur le chemin

from pptx import Presentation  # noqa: E402

from app.services import translation_runner as tr  # noqa: E402
from app.services.jobs import jobs  # noqa: E402

TEXTES = ["Premier texte", "Deuxieme texte", "Troisieme texte"]


def _pptx(chemin: str) -> bytes:
    prs = Presentation()
    for texte in TEXTES:
        diapo = prs.slides.add_slide(prs.slide_layouts[5])
        diapo.shapes.title.text = texte
    prs.save(chemin)
    with open(chemin, "rb") as f:
        return f.read()


def _cible(texte: str) -> str:
    """Traduction factice, DANS les balises de run : « [[0]]x[[/0]] » → « [[0]]T:x[[/0]] »."""
    return texte.replace("]]", "]]T:", 1)


def _nu(texte: str) -> str:
    return texte.replace("[[0]]", "").replace("[[/0]]", "")


class Traducteur:
    """Préfixe chaque fragment ; échoue sur `echoue_sur` tant qu'il est posé."""

    def __init__(self):
        self.appels: list[str] = []
        self.echoue_sur: str | None = None

    def translate_json(self, chemin, target_lang, progress_callback=None, **_):
        with open(chemin, encoding="utf-8") as f:
            data = json.load(f)
        sources = [_nu(el["text"]) for sd in data["slides"] for el in tr._elements_de(sd)]
        self.appels.extend(sources)
        if self.echoue_sur and any(self.echoue_sur in s for s in sources):
            return False, "fournisseur coupé"
        for sd in data["slides"]:
            for el in tr._elements_de(sd):
                el["translated_text"] = _cible(el["text"])
        sortie = chemin + ".out.json"
        with open(sortie, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False)
        return True, sortie

    def _translate_batch(self, *a, **k):      # passe de cohérence : rien à faire
        return None


def _lancer(dossier: str, octets: bytes, nom: str) -> tuple[str, dict]:
    job_id = jobs.create()
    original = os.path.join(dossier, "original.pptx")
    traduction = os.path.join(dossier, "translated.json")
    sortie = os.path.join(dossier, f"render_{nom}.pptx")
    tr.run_pptx_progressive_job(
        job_id, octets, original, sortie, f"{nom}.pptx",
        os.path.join(dossier, f"partial_{nom}.pdf"), traduction, "en")
    return traduction, jobs.get(job_id) or {}


def _textes_sortie(chemin: str) -> list[str]:
    prs = Presentation(chemin)
    return [sh.text_frame.text for s in prs.slides for sh in s.shapes
            if sh.has_text_frame and sh.text_frame.text]


def main() -> int:
    checks: list[tuple[str, bool, str]] = []

    def ok(nom, cond, detail=""):
        checks.append((nom, bool(cond), detail))

    def _sans_libreoffice(*_a, **_k):
        raise RuntimeError("LibreOffice neutralisé")

    tr.convert_to_pdf_bytes = _sans_libreoffice
    trad = Traducteur()
    tr.ai_translator, tr.ai_active = trad, True

    dossier = tempfile.mkdtemp(prefix="reprise_pptx_")
    try:
        octets = _pptx(os.path.join(dossier, "source.pptx"))
        traduction = os.path.join(dossier, "translated.json")
        reprise = tr._chemin_reprise(traduction)

        # ── 1. Premier essai : la diapo 3 échoue ─────────────────────────
        trad.echoue_sur = "Troisieme"
        _lancer(dossier, octets, "a")
        ok("1. échec : aucune traduction finale", not os.path.exists(traduction))
        gardees = tr._lire_reprise(reprise)
        ok("1. échec : diapos 1 et 2 gardées", sorted(gardees) == [1, 2],
           str(sorted(gardees)))

        # ── 2. Relance : seule la diapo manquante part au modèle ─────────
        trad.echoue_sur, trad.appels = None, []
        _lancer(dossier, octets, "b")
        # Rien de ce qui était déjà traduit ne repart au modèle, et la diapo
        # en échec, elle, y part bien.
        ok("2. relance : seuls les fragments manquants retraduits",
           "Troisieme texte" in trad.appels
           and not {"Premier texte", "Deuxieme texte"} & set(trad.appels)
           and len(trad.appels) == len(set(trad.appels)), str(trad.appels))
        # Le gabarit (texte « Click to edit… ») est pris par une diapo ou une
        # autre selon la course : il doit être repris quel que soit son porteur.
        ok("2. relance : traduction finale écrite", os.path.isfile(traduction))
        if os.path.isfile(traduction):
            with open(traduction, encoding="utf-8") as f:
                final = json.load(f)
            paires = [(_nu(el["text"]), el.get("translated_text"))
                      for sd in final["slides"] for el in tr._elements_de(sd)]
            ok("2. relance : tout traduit, source conservée",
               paires and all(c == _cible(f"[[0]]{s}[[/0]]") for s, c in paires)
               and {s for s, _ in paires} >= set(TEXTES), str(paires))
        ok("2. relance : fichier de reprise retiré", not os.path.exists(reprise))
        sortie = os.path.join(dossier, "render_b.pptx")
        ok("2. relance : PPTX final traduit",
           os.path.isfile(sortie) and _textes_sortie(sortie) == ["T:" + t for t in TEXTES],
           str(_textes_sortie(sortie)) if os.path.isfile(sortie) else "absent")

        # ── 3. Reprise périmée : un texte source qui ne correspond plus ──
        os.remove(traduction)
        perimee = {1: {"text_elements": [{"id": "x", "text": "autre chose",
                                          "translated_text": "FAUX"}]}}
        tr._ecrire_reprise(reprise, perimee)
        trad.appels = []
        _lancer(dossier, octets, "c")
        ok("3. reprise périmée : rien n'est repris à tort",
           set(TEXTES) <= set(trad.appels), str(trad.appels))

        # ── 4. Original à 0 octet : réécrit, pas réutilisé ───────────────
        original = os.path.join(dossier, "original.pptx")
        open(original, "wb").close()
        tr._sauver_original(original, octets)
        ok("4. original vide réécrit", os.path.getsize(original) == len(octets))
        tr._sauver_original(original, b"")
        ok("4. original complet gardé sans octets neufs",
           os.path.getsize(original) == len(octets))
        open(original, "wb").close()
        try:
            tr._sauver_original(original, b"")
            ok("4. original vide sans octets neufs refusé", False)
        except ValueError:
            ok("4. original vide sans octets neufs refusé", True)
    finally:
        shutil.rmtree(dossier, ignore_errors=True)

    for nom, passe, detail in checks:
        print(("OK   " if passe else "ECHEC"), nom, ("" if passe else detail))
    n_ok = sum(1 for _, p, _ in checks if p)
    print(f"{n_ok}/{len(checks)}")
    return 0 if n_ok == len(checks) else 1


if __name__ == "__main__":
    sys.exit(main())
